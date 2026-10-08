import os
import re
import unicodedata
from pathlib import Path

import pandas as pd

from config import DATASET_PATH as ABSTRACTS_DIR
from config import SOURCE_LIST_PATH
from utils.areas import SUBJECT_AREAS

FIRST_YEAR = 2019
LAST_YEAR = 2026
RAW_COLUMNS = {
    "EID": "eid",
    "DOI": "doi",
    "Title": "title",
    "Abstract": "abstract",
    "Year": "year",
    "Source title": "source_title",
    "Volume": "volume",
    "Issue": "issue",
    "Art. No.": "article_no",
    "Page start": "page_start",
}
EM_DASH = "—"
NO_ABSTRACT = "[No abstract available]"
CORRECTED_EXPORT_DIRS = ("correction-a", "correction-y")


def normalize_journal_name(name):
    if pd.isna(name):
        return pd.NA
    name = unicodedata.normalize("NFKD", str(name))
    name = name.encode("ascii", "ignore").decode("ascii").lower()
    name = re.sub("\\([^)]*\\)", " ", name)
    name = name.replace("&", " and ")
    name = re.sub("[^a-z0-9\\s]", " ", name)
    return re.sub("\\s+", " ", name).strip()


def extract_subarea(highest_percentile):
    if pd.isna(highest_percentile):
        return pd.NA
    subarea = str(highest_percentile).strip().splitlines()[-1].strip()
    return "Multidisciplinary" if subarea.endswith("/223") else subarea


def export_paths(directory=ABSTRACTS_DIR):

    def subdirectory_order(parent, subdirectory):
        is_top_level = parent == directory
        priority = (
            CORRECTED_EXPORT_DIRS.index(subdirectory)
            if is_top_level and subdirectory in CORRECTED_EXPORT_DIRS
            else len(CORRECTED_EXPORT_DIRS)
        )
        return (priority, subdirectory)

    def paths_with_suffix(suffix):
        paths = []
        for parent, subdirectories, filenames in os.walk(directory):
            parent = Path(parent)
            subdirectories.sort(key=lambda name: subdirectory_order(parent, name))
            paths += [
                parent / name for name in sorted(filenames) if name.endswith(suffix)
            ]
        return paths

    return paths_with_suffix(".csv") + paths_with_suffix(".xlsx")


def load_raw_exports(directory=ABSTRACTS_DIR, chunksize=10000):
    """Read exports in batches, keeping the first record for each EID."""
    paths = export_paths(directory)
    if not paths:
        raise FileNotFoundError(f"No Scopus exports under {directory}")
    exports, seen, n_read = [], set(), 0
    for path in paths:
        if path.suffix == ".csv":
            batches = pd.read_csv(
                path, encoding="utf-8-sig", dtype=str,
                usecols=lambda name: name in RAW_COLUMNS, chunksize=chunksize,
            )
        else:
            batches = excel_batches(path, chunksize)
        for export in batches:
            export = export.reindex(columns=list(RAW_COLUMNS))
            n_read += len(export)
            export = export.dropna(subset=["EID"]).drop_duplicates("EID")
            export = export[~export.EID.isin(seen)]
            seen.update(export.EID)
            exports.append(export)
    raw_records = pd.concat(exports, ignore_index=True).rename(columns=RAW_COLUMNS)
    raw_records["year"] = pd.to_numeric(raw_records["year"], errors="coerce")
    print(f"Export files read:      {len(paths)}")
    print(f"Records read:           {n_read:,}")
    raw_records = raw_records.dropna(subset=["eid"]).drop_duplicates(subset="eid")
    print(f"Records after dedup:    {len(raw_records):,}")
    return raw_records.reset_index(drop=True)


def excel_batches(path, chunksize):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = next(rows)
        columns = [(i, name) for i, name in enumerate(header) if name in RAW_COLUMNS]
        batch = []
        for row in rows:
            batch.append({name: None if row[i] is None else str(row[i]) for i, name in columns})
            if len(batch) == chunksize:
                yield pd.DataFrame(batch)
                batch = []
        if batch:
            yield pd.DataFrame(batch)
    finally:
        workbook.close()


def has_abstract(records):
    text = records["abstract"].fillna("").str.strip()
    return text.ne("") & text.ne(NO_ABSTRACT)


def load_journal_areas(path=SOURCE_LIST_PATH):
    subarea_to_area = {
        subarea: area
        for area, subareas in SUBJECT_AREAS.items()
        for subarea in subareas
    }
    source_list = pd.read_csv(path, encoding="utf-8-sig")
    journals = source_list[["Source title", "Highest percentile"]].copy()
    journals["journal"] = journals["Source title"].apply(normalize_journal_name)
    journals["subarea"] = journals["Highest percentile"].apply(extract_subarea)
    journals["area"] = journals["subarea"].map(subarea_to_area)
    journals.loc[journals["subarea"].eq("Multidisciplinary"), "area"] = (
        "Multidisciplinary"
    )
    journals = journals.dropna(subset=["subarea", "area"])
    journals = (
        journals.groupby("journal", group_keys=False)
        .filter(lambda group: len(group[["subarea", "area"]].drop_duplicates()) == 1)
        .drop_duplicates("journal")
    )
    print(f"Journals classified:    {len(journals):,}")
    return journals[["journal", "subarea", "area"]]


def is_erratum(records):
    starts_with_erratum = records["abstract"].fillna("").str.strip().str.startswith(
        ("Erratum", "Correction")
    ) | records["title"].fillna("").str.strip().str.startswith(
        ("Erratum", "Correction")
    )
    return starts_with_erratum


def build_corpus(raw_records, journals):
    corpus = (
        raw_records.assign(
            journal=raw_records["source_title"].apply(normalize_journal_name)
        )
        .merge(journals, on="journal", how="left", validate="many_to_one")
        .dropna(subset=["area"])
    )
    print(f"In a classified journal: {len(corpus):,}")
    corpus = corpus[corpus["year"].between(FIRST_YEAR, LAST_YEAR)]
    print(f"Published {FIRST_YEAR}-{LAST_YEAR}:    {len(corpus):,}")
    with_abstract = has_abstract(corpus)
    print(f"No abstract, dropped:   {(~with_abstract).sum():,}")
    corpus = corpus[with_abstract]
    erratum = is_erratum(corpus)
    print(f"Errata, dropped:        {erratum.sum():,}")
    corpus = corpus[~erratum]
    print(
        f"Final corpus:           {len(corpus):,} abstracts from {corpus['journal'].nunique():,} journals"
    )
    return corpus.reset_index(drop=True)


def em_dash_rule_variants(abstracts):
    other_em_width_dashes = "[—―⸺⸻]"
    html_entities = "&mdash;|&#8212;|&#x2014;"
    mojibake = "â€”"
    repaired = abstracts.str.replace(html_entities, EM_DASH, regex=True).str.replace(
        mojibake, EM_DASH, regex=False
    )
    without_numeric_ranges = abstracts.str.replace(
        "(?<=\\d)\\s?" + EM_DASH + "\\s?(?=\\d)", " ", regex=True
    )
    variants = {
        "literal U+2014 (primary)": abstracts.str.contains(EM_DASH, regex=False),
        "any em-width dash": abstracts.str.contains(other_em_width_dashes, regex=True),
        "entities and mojibake repaired": repaired.str.contains(EM_DASH, regex=False),
        "numeric ranges excluded": without_numeric_ranges.str.contains(
            EM_DASH, regex=False
        ),
    }
    return pd.DataFrame(
        {
            "rule": list(variants),
            "n_with_em_dash": [int(hits.sum()) for hits in variants.values()],
            "pct": [100 * hits.mean() for hits in variants.values()],
        }
    )
